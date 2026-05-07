"""
Genera los documentos Word de documentación técnica y manual de usuario.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

# ──────────────────────────────────────────────
# Utilidades
# ──────────────────────────────────────────────

AZUL_OSCURO  = RGBColor(0x1E, 0x3A, 0x5F)  # #1E3A5F
AZUL_MEDIO   = RGBColor(0x25, 0x63, 0xEB)  # #2563EB
GRIS_TEXTO   = RGBColor(0x37, 0x41, 0x51)  # #374151
GRIS_CLARO   = RGBColor(0x6B, 0x72, 0x80)  # #6B7280
BLANCO       = RGBColor(0xFF, 0xFF, 0xFF)
FONDO_TABLA  = RGBColor(0xEF, 0xF6, 0xFF)  # #EFF6FF
FONDO_HEADER = RGBColor(0x1E, 0x3A, 0x5F)  # azul oscuro


def set_cell_bg(cell, color: RGBColor):
    """Aplica color de fondo a una celda."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    hex_color = str(color)  # RGBColor.__str__ devuelve el hex sin '#'
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_border(cell, top=None, bottom=None, left=None, right=None):
    """Agrega bordes a una celda."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        if val:
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), val)
            tcBorders.append(el)
    tcPr.append(tcBorders)


def heading(doc, text, level=1, color=AZUL_OSCURO, space_before=18, space_after=6):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    for run in p.runs:
        run.font.color.rgb = color
        if level == 1:
            run.font.size = Pt(18)
            run.font.bold = True
        elif level == 2:
            run.font.size = Pt(14)
            run.font.bold = True
        elif level == 3:
            run.font.size = Pt(12)
            run.font.bold = True
    return p


def body(doc, text, color=GRIS_TEXTO, size=11, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.size  = Pt(size)
    run.font.color.rgb = color
    return p


def bullet(doc, text, bold_prefix=None, color=GRIS_TEXTO):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        r1.font.color.rgb = AZUL_MEDIO
        r1.font.size = Pt(11)
    r2 = p.add_run(text)
    r2.font.size  = Pt(11)
    r2.font.color.rgb = color
    return p


def code_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent  = Cm(1)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    # fondo gris muy claro simulado con sombreado
    rPr = run._r.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F3F4F6")
    rPr.append(shd)
    return p


def divider(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"),   "single")
    bottom.set(qn("w:sz"),    "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "CBD5E1")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def styled_table(doc, headers, rows, col_widths=None):
    """Crea una tabla estilizada con encabezado azul."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    # Encabezado
    hdr_row = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr_row.cells[i]
        set_cell_bg(cell, FONDO_HEADER)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.bold = True
        run.font.color.rgb = BLANCO
        run.font.size = Pt(10)

    # Filas de datos
    for r_idx, row_data in enumerate(rows):
        tr = table.rows[r_idx + 1]
        bg = FONDO_TABLA if r_idx % 2 == 0 else RGBColor(0xFF, 0xFF, 0xFF)
        for c_idx, cell_text in enumerate(row_data):
            cell = tr.cells[c_idx]
            set_cell_bg(cell, bg)
            p = cell.paragraphs[0]
            run = p.add_run(str(cell_text))
            run.font.size = Pt(10)
            run.font.color.rgb = GRIS_TEXTO

    # Anchos de columna
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)

    doc.add_paragraph()  # espacio después
    return table


def cover_page(doc, title, subtitle, team, year="2026"):
    """Crea una portada profesional."""
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph()

    # Título
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size  = Pt(28)
    run.font.color.rgb = AZUL_OSCURO
    p.paragraph_format.space_after = Pt(12)

    # Subtítulo
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run(subtitle)
    run2.font.size  = Pt(16)
    run2.font.color.rgb = AZUL_MEDIO
    p2.paragraph_format.space_after = Pt(40)

    divider(doc)

    doc.add_paragraph()

    # Equipo
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run("Equipo de trabajo")
    r3.bold = True
    r3.font.size  = Pt(12)
    r3.font.color.rgb = GRIS_CLARO
    p3.paragraph_format.space_after = Pt(6)

    for nombre in team:
        p4 = doc.add_paragraph()
        p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r4 = p4.add_run(nombre)
        r4.font.size  = Pt(12)
        r4.font.color.rgb = AZUL_OSCURO
        p4.paragraph_format.space_after = Pt(4)

    doc.add_paragraph()
    divider(doc)
    doc.add_paragraph()

    # Curso y año
    for line in ["Gestión del Conocimiento", f"Proyecto Educativo · {year}"]:
        p5 = doc.add_paragraph()
        p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r5 = p5.add_run(line)
        r5.font.size  = Pt(11)
        r5.font.color.rgb = GRIS_CLARO
        p5.paragraph_format.space_after = Pt(4)

    doc.add_page_break()


def configure_doc(doc):
    """Configura márgenes y fuente por defecto."""
    section = doc.sections[0]
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3)
    section.right_margin  = Cm(2.5)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)


# ══════════════════════════════════════════════
# DOCUMENTO 1: DOCUMENTACIÓN TÉCNICA
# ══════════════════════════════════════════════

def crear_doc_tecnico():
    doc = Document()
    configure_doc(doc)

    team = [
        "Luis Daniel Martinez Rincon",
        "Horacio Palacios Palacios",
        "Dangelly Angelica Blandon Negrete",
        "Karime Salazar Hinestroza",
    ]

    cover_page(doc,
               "Documentación Técnica",
               "Sitio Web — Gestión del Conocimiento",
               team)

    # ── 1. Descripción General ──────────────────
    heading(doc, "1. Descripción General")
    body(doc,
         "Sitio web educativo de una sola página (Single Page Application) que presenta los fundamentos "
         "teóricos y prácticos de la Gestión del Conocimiento. Cubre cuatro subtemas principales con "
         "contenido académico, videos embebidos, estadísticas y bibliografía completa con acceso directo "
         "a las fuentes.")
    body(doc, "URL de producción: https://gestion-conocimiento.vercel.app", color=AZUL_MEDIO)
    divider(doc)

    # ── 2. Stack Tecnológico ────────────────────
    heading(doc, "2. Stack Tecnológico")
    styled_table(doc,
        ["Capa", "Tecnología", "Versión"],
        [
            ["Framework",   "Next.js (Pages Router)",   "^16.2.4"],
            ["UI",          "React",                    "^18.3.1"],
            ["Estilos",     "Tailwind CSS",             "^3.4.3"],
            ["Animaciones", "Framer Motion",            "^11.1.9"],
            ["SEO",         "next-seo",                 "^6.5.0"],
            ["Íconos",      "react-icons",              "^5.1.0"],
            ["Linting",     "ESLint + eslint-config-next", "^8.57.0"],
            ["Despliegue",  "Vercel",                   "—"],
        ],
        col_widths=[4, 8, 4]
    )
    divider(doc)

    # ── 3. Estructura de Archivos ───────────────
    heading(doc, "3. Estructura de Archivos")
    tree = (
        "Trabajo Final/\n"
        "├── pages/\n"
        "│   ├── index.js          Página principal\n"
        "│   ├── _app.js           Wrapper global\n"
        "│   ├── _document.js      HTML base\n"
        "│   └── sitemap.xml.js    Sitemap dinámico\n"
        "├── components/\n"
        "│   ├── SEO.jsx           Meta tags y datos estructurados\n"
        "│   ├── Navbar.jsx        Barra de navegación fija\n"
        "│   ├── Hero.jsx          Sección portada\n"
        "│   ├── TopicSection.jsx  Sección reutilizable por subtema\n"
        "│   ├── ResumenSection.jsx Conceptos transversales\n"
        "│   ├── FuentesSection.jsx Bibliografía filtrable\n"
        "│   └── Footer.jsx        Pie de página\n"
        "├── data/\n"
        "│   └── content.js        FUENTE ÚNICA DE DATOS\n"
        "├── styles/\n"
        "│   └── globals.css       Estilos globales + Tailwind\n"
        "├── public/\n"
        "│   └── robots.txt\n"
        "├── docs/\n"
        "│   ├── documentacion-tecnica.docx\n"
        "│   └── manual-de-usuario.docx\n"
        "├── next.config.js\n"
        "├── tailwind.config.js\n"
        "└── package.json"
    )
    code_block(doc, tree)
    divider(doc)

    # ── 4. Arquitectura ─────────────────────────
    heading(doc, "4. Arquitectura de la Aplicación")
    body(doc, "El flujo de datos es unidireccional: data/content.js → pages/index.js → Componentes. "
              "Ningún componente genera ni modifica datos propios.")
    arch = (
        "pages/index.js\n"
        "├── <SEO />             Meta tags dinámicos\n"
        "├── <Navbar />          Navegación fija global\n"
        "├── <main>\n"
        "│   ├── <Hero />                         Portada\n"
        "│   ├── <TopicSection tema={tema1} />    Subtema 1\n"
        "│   ├── <TopicSection tema={tema2} />    Subtema 2\n"
        "│   ├── <TopicSection tema={tema3} />    Subtema 3\n"
        "│   ├── <TopicSection tema={tema4} />    Subtema 4\n"
        "│   ├── <ResumenSection />               Síntesis\n"
        "│   └── <FuentesSection fuentes={...} /> Bibliografía\n"
        "└── <Footer />"
    )
    code_block(doc, arch)
    divider(doc)

    # ── 5. Estructura de Datos ──────────────────
    heading(doc, "5. Estructura de Datos (data/content.js)")

    heading(doc, "5.1 siteConfig", level=2)
    fields_config = [
        ["title",       "string",   "Título del sitio (SEO y Open Graph)"],
        ["description", "string",   "Descripción para motores de búsqueda"],
        ["url",         "string",   "URL de producción"],
        ["author",      "string",   "Crédito del proyecto"],
        ["keywords",    "string[]", "Palabras clave para SEO"],
        ["team",        "string[]", "Nombres de los integrantes"],
    ]
    styled_table(doc, ["Campo", "Tipo", "Descripción"], fields_config, col_widths=[4, 3, 9])

    heading(doc, "5.2 Objeto tema (tema1 – tema4)", level=2)
    fields_tema = [
        ["id",           "string",   "Identificador único (usado en id HTML y SEO)"],
        ["titulo",       "string",   "Título del subtema"],
        ["icon",         "string",   "Emoji decorativo"],
        ["color",        "string",   "Paleta: blue | amber | green | purple"],
        ["resumen",      "string",   "Párrafo introductorio del header"],
        ["estadisticas", "array",    "Tarjetas de métricas con valor y descripción"],
        ["contenido",    "array",    "Bloques de texto: subtitulo + texto (6 por tema)"],
        ["videos",       "array",    "Videos de YouTube: title, id, descripcion (3 por tema)"],
    ]
    styled_table(doc, ["Campo", "Tipo", "Descripción"], fields_tema, col_widths=[4, 3, 9])

    heading(doc, "5.3 Objeto fuente (array fuentes)", level=2)
    fields_fuente = [
        ["id",        "number",  "Número correlativo de referencia"],
        ["tipo",      "string",  "Libro | Artículo | Informe"],
        ["temas",     "number[]","Subtemas relacionados (1, 2, 3 o 4)"],
        ["titulo",    "string",  "Título completo de la obra"],
        ["autores",   "string",  "Autores en formato APA"],
        ["año",       "number",  "Año de publicación"],
        ["doi",       "string",  "DOI — genera botón verde 'Ver artículo'"],
        ["url",       "string",  "URL directa — genera botón azul 'Ver fuente'"],
        ["isbn",      "string",  "ISBN-13 (libros)"],
        ["editorial", "string",  "Editorial (libros)"],
    ]
    styled_table(doc, ["Campo", "Tipo", "Descripción"], fields_fuente, col_widths=[4, 3, 9])
    divider(doc)

    # ── 6. Sistema de Colores ───────────────────
    heading(doc, "6. Sistema de Colores")
    styled_table(doc,
        ["Color (campo color)", "Subtema", "Base Tailwind"],
        [
            ["blue",   "Explícito vs Tácito",           "primary-*"],
            ["amber",  "Innovación y Cultura",           "amber-*"],
            ["green",  "Trabajo Colaborativo",           "emerald-*"],
            ["purple", "Aprendizaje Organizacional",     "purple-*"],
        ],
        col_widths=[5, 7, 4]
    )
    divider(doc)

    # ── 7. Componentes ──────────────────────────
    heading(doc, "7. Componentes — Referencia")

    componentes = [
        ("Navbar.jsx", [
            "Barra de navegación fija (position: fixed, z-50).",
            "Cambia de transparente a bg-white/95 con backdrop-blur al superar 60 px de scroll.",
            "Menú hamburger en móvil; se cierra al hacer clic fuera.",
            "Links: Inicio, 4 subtemas, Fuentes.",
        ]),
        ("Hero.jsx", [
            "Portada con gradiente from-primary-900 via-primary-800 to-primary-600.",
            "Muestra nombres del equipo desde siteConfig.team.",
            "Dos CTAs: 'Explorar contenido' (#subtema-1) y 'Ver fuentes' (#fuentes).",
            "Flecha de scroll animada al fondo.",
        ]),
        ("TopicSection.jsx", [
            "Componente reutilizable parametrizado con tema e index.",
            "Índice par = fondo blanco; índice impar = fondo de color suave.",
            "Sub-secciones: header, estadísticas (grid 3 col), contenido (grid 2 col), videos.",
            "Iframes con youtube-nocookie.com (modo privacidad reforzada).",
            "Aspecto 16:9 con paddingTop: 56.25% + hijo position: absolute.",
        ]),
        ("FuentesSection.jsx", [
            "Filtros interactivos por subtema con useState.",
            "Botón verde 'Ver artículo (DOI)' cuando existe f.doi.",
            "Botón azul 'Ver fuente' cuando existe f.url.",
            "Nota metodológica sobre formato APA 7.ª edición.",
        ]),
        ("Footer.jsx", [
            "Grid de 4 columnas: Marca, Subtemas, Recursos, Integrantes.",
            "Equipo generado dinámicamente desde siteConfig.team.",
        ]),
    ]

    for nombre, items in componentes:
        heading(doc, nombre, level=2, space_before=10)
        for item in items:
            bullet(doc, item)

    divider(doc)

    # ── 8. Scripts ──────────────────────────────
    heading(doc, "8. Scripts de Desarrollo")
    scripts = [
        ("npm run dev",   "Inicia servidor de desarrollo (puerto 3000)"),
        ("npm run build", "Construye la aplicación para producción"),
        ("npm start",     "Inicia servidor de producción (requiere build)"),
        ("npm run lint",  "Analiza el código con ESLint"),
    ]
    styled_table(doc, ["Comando", "Descripción"], scripts, col_widths=[6, 10])
    divider(doc)

    # ── 9. Despliegue ───────────────────────────
    heading(doc, "9. Despliegue en Vercel")
    body(doc, "El proyecto se despliega automáticamente al hacer push a la rama main. "
              "No requiere variables de entorno. El proceso es:")
    for paso in [
        "Push a main → Vercel detecta cambios automáticamente.",
        "Vercel ejecuta npm run build.",
        "Next.js genera las páginas estáticas.",
        "Vercel despliega en CDN global sin intervención manual.",
    ]:
        bullet(doc, paso)
    divider(doc)

    # ── 10. Cómo Modificar el Contenido ─────────
    heading(doc, "10. Cómo Modificar el Contenido")
    body(doc, "Todo el contenido del sitio vive en data/content.js. Para actualizar el sitio:", space_after=4)

    modificaciones = [
        ("Cambiar texto de un subtema: ",
         "Editar los campos titulo, resumen o los objetos dentro de contenido[] en tema1–tema4."),
        ("Agregar un bloque de contenido: ",
         "Añadir { subtitulo: '...', texto: '...' } al array contenido del tema."),
        ("Cambiar un video: ",
         "Reemplazar el campo id con el ID de YouTube (parte ?v=XXXXXXX de la URL) y actualizar title y descripcion."),
        ("Agregar una fuente: ",
         "Añadir un objeto al array fuentes con id correlativo, tipo, temas y los campos del tipo (Libro, Artículo o Informe)."),
        ("Cambiar los integrantes: ",
         "Editar el array siteConfig.team. Los cambios se reflejan en Hero y Footer automáticamente."),
    ]
    for prefix, text in modificaciones:
        bullet(doc, text, bold_prefix=prefix)

    divider(doc)

    # ── 11. Decisiones de Diseño ─────────────────
    heading(doc, "11. Decisiones de Diseño")
    decisiones = [
        ["Single Page Application", "Favorece la lectura continua y el scroll narrativo"],
        ["Datos en content.js",     "Un solo punto de modificación para todo el contenido"],
        ["youtube-nocookie.com",    "Reduce cookies de seguimiento para los visitantes"],
        ["Open Library para libros","URLs canónicas y permanentes basadas en ISBN"],
        ["DOI para artículos",      "Identificador permanente internacional"],
        ["Tailwind utility-first",  "Estilos colocalizados, sin archivos CSS externos"],
        ["Paleta por subtema",      "Diferenciación visual inmediata entre secciones"],
    ]
    styled_table(doc, ["Decisión", "Justificación"], decisiones, col_widths=[7, 9])

    # Pie
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("Documentación técnica — versión 1.0 · mayo 2026")
    run.font.size  = Pt(9)
    run.font.color.rgb = GRIS_CLARO
    run.italic = True

    doc.save(r"d:\Archivos\Escritorio\GESTION DEL CONOCIMIENTO\Trabajo Final\docs\documentacion-tecnica.docx")
    print("OK: documentacion-tecnica.docx generado")


# ══════════════════════════════════════════════
# DOCUMENTO 2: MANUAL DE USUARIO
# ══════════════════════════════════════════════

def crear_manual_usuario():
    doc = Document()
    configure_doc(doc)

    team = [
        "Luis Daniel Martinez Rincon",
        "Horacio Palacios Palacios",
        "Dangelly Angelica Blandon Negrete",
        "Karime Salazar Hinestroza",
    ]

    cover_page(doc,
               "Manual de Usuario",
               "Sitio Web — Gestión del Conocimiento",
               team)

    # Intro
    heading(doc, "¿Qué es este sitio?")
    body(doc,
         "Este sitio web es un recurso educativo interactivo sobre Gestión del Conocimiento. "
         "Presenta de forma clara los conceptos clave de esta disciplina con contenido académico, "
         "videos explicativos embebidos y bibliografía completa con acceso directo a las fuentes originales.")

    heading(doc, "Temas que cubre:", level=3, space_before=6)
    for t in ["Conocimiento explícito y tácito en las organizaciones",
              "Innovación y cultura organizacional",
              "Trabajo colaborativo y comunidades de práctica",
              "Aprendizaje organizacional y memoria colectiva"]:
        bullet(doc, t)
    divider(doc)

    # 1. Acceso
    heading(doc, "1. Acceso al Sitio")
    body(doc, "Ingresa desde cualquier navegador moderno a:")
    p = doc.add_paragraph()
    run = p.add_run("https://gestion-conocimiento.vercel.app")
    run.bold = True
    run.font.size  = Pt(13)
    run.font.color.rgb = AZUL_MEDIO
    p.paragraph_format.space_after = Pt(8)

    body(doc, "Compatible con:")
    for b in ["Google Chrome (recomendado)", "Mozilla Firefox", "Microsoft Edge",
              "Safari (Mac / iOS)", "Navegadores móviles Android e iOS"]:
        bullet(doc, b)
    body(doc, "No requiere registro ni instalación de ningún software adicional.", color=GRIS_CLARO)
    divider(doc)

    # 2. Estructura
    heading(doc, "2. Estructura General del Sitio")
    body(doc, "El sitio es una sola página larga dividida en secciones. Puedes navegar "
              "entre ellas usando el menú superior o haciendo scroll hacia abajo.")
    styled_table(doc,
        ["Sección", "Descripción"],
        [
            ["Inicio",                    "Portada con presentación general y accesos rápidos"],
            ["Subtema 1",                 "Conocimiento Explícito vs Tácito"],
            ["Subtema 2",                 "Innovación y Cultura Organizacional"],
            ["Subtema 3",                 "Trabajo Colaborativo y Comunidades de Práctica"],
            ["Subtema 4",                 "Aprendizaje Organizacional y Memoria Colectiva"],
            ["Conceptos Transversales",   "Síntesis e hitos históricos de la disciplina"],
            ["Fuentes",                   "Bibliografía completa con acceso a las fuentes"],
        ],
        col_widths=[6, 10]
    )
    divider(doc)

    # 3. Navbar
    heading(doc, "3. La Barra de Navegación")
    body(doc, "La barra de navegación está siempre visible en la parte superior de la pantalla "
              "(barra fija). Adapta su apariencia según el dispositivo:")

    heading(doc, "En pantallas grandes (computador):", level=3, space_before=8)
    for item in [
        "Todos los enlaces del menú se muestran de izquierda a derecha.",
        "Al inicio es transparente sobre el fondo oscuro de la portada.",
        "Al hacer scroll, se vuelve blanca con sombra suave para mejor legibilidad.",
    ]:
        bullet(doc, item)

    heading(doc, "En pantallas pequeñas (celular / tablet):", level=3, space_before=8)
    for item in [
        "Verás el logo y un ícono de tres líneas (☰) en la esquina derecha.",
        "Toca el ícono para abrir el menú desplegable.",
        "El menú se cierra automáticamente al seleccionar una opción.",
    ]:
        bullet(doc, item)

    body(doc, "El ícono 🧩 (rompecabezas) en la esquina superior izquierda siempre te lleva de vuelta al inicio.")
    divider(doc)

    # 4. Portada
    heading(doc, "4. Sección de Inicio (Portada)")
    body(doc, "Al abrir el sitio verás la portada con los siguientes elementos:")
    for item in [
        ("Título principal: ", "\"Gestión del Conocimiento\" con acento visual en dorado."),
        ("Descripción: ",       "Párrafo introductorio sobre el alcance del sitio."),
        ("Equipo: ",            "Nombres de los integrantes listados debajo de la descripción."),
        ("\"Explorar contenido\": ", "Botón que te lleva directamente al Subtema 1."),
        ("\"Ver fuentes\": ",   "Botón que te lleva a la sección de bibliografía."),
        ("Tarjetas de temas: ", "Cuatro tarjetas con íconos que resumen los temas del sitio."),
        ("Flecha animada: ",    "En la parte inferior, indica que puedes seguir haciendo scroll."),
    ]:
        bullet(doc, item[1], bold_prefix=item[0])
    divider(doc)

    # 5. Subtemas
    heading(doc, "5. Secciones de Subtemas")
    body(doc, "Cada uno de los cuatro subtemas tiene la misma estructura. "
              "Reconocerás cada sección por su color característico:")
    styled_table(doc,
        ["Subtema", "Color", "Ícono"],
        [
            ["Explícito vs Tácito",           "Azul",             "🧠"],
            ["Innovación y Cultura",           "Ámbar / Naranja",  "💡"],
            ["Trabajo Colaborativo",           "Verde esmeralda",  "🤝"],
            ["Aprendizaje Organizacional",     "Púrpura",          "📚"],
        ],
        col_widths=[7, 5, 3]
    )

    heading(doc, "5.1 Encabezado del Subtema", level=2)
    for item in [
        "Un ícono representativo en una caja redondeada.",
        "Una etiqueta que indica el número de subtema (ej. \"Subtema 1\").",
        "El título completo del subtema.",
        "Un párrafo resumen que presenta los conceptos clave de la sección.",
    ]:
        bullet(doc, item)

    heading(doc, "5.2 Estadísticas", level=2)
    body(doc, "Tres tarjetas de color sólido muestran datos cuantitativos relevantes al tema "
              "(porcentajes, cifras de estudios académicos). Son puramente informativas.")

    heading(doc, "5.3 Bloques de Contenido", level=2)
    body(doc, "Seis tarjetas blancas con borde de color presentan el contenido académico. Cada tarjeta tiene:")
    for item in [
        "Un punto de color a la izquierda como decoración.",
        "Un subtítulo en negrita.",
        "Un párrafo con la explicación del concepto e inclusión de referencias académicas.",
    ]:
        bullet(doc, item)
    body(doc, "Pasa el cursor sobre las tarjetas para ver el efecto de elevación (sombra).", color=GRIS_CLARO)

    heading(doc, "5.4 Videos Relacionados", level=2)
    body(doc, "Al final de cada subtema encontrarás tres videos de YouTube embebidos directamente "
              "en la página sobre un fondo oscuro. Cómo usarlos:")
    for item in [
        "Los videos cargan de forma diferida (lazy loading) para no ralentizar la página.",
        "Haz clic en el botón ▶ para reproducir directamente en el sitio.",
        "Usa los controles del reproductor para pausar, avanzar o ajustar el volumen.",
        "El botón ⛶ en la esquina inferior derecha amplía el video a pantalla completa.",
        "Si prefieres verlo en YouTube, haz clic en el enlace \"Ver en YouTube\" debajo del video.",
    ]:
        bullet(doc, item)
    body(doc,
         "Nota de privacidad: los videos se cargan desde youtube-nocookie.com, "
         "el dominio de privacidad reforzada de YouTube que reduce el seguimiento por cookies.",
         color=GRIS_CLARO)
    divider(doc)

    # 6. Conceptos Transversales
    heading(doc, "6. Sección de Conceptos Transversales")
    body(doc, "Entre los subtemas y la bibliografía encontrarás una sección de síntesis con:")
    bullet(doc, "Los conceptos que conectan todos los subtemas entre sí.")
    bullet(doc, "Una línea de tiempo con hitos históricos clave de la Gestión del Conocimiento.")
    body(doc, "Esta sección es de lectura; no tiene interacciones.", color=GRIS_CLARO)
    divider(doc)

    # 7. Fuentes
    heading(doc, "7. Sección de Fuentes Bibliográficas")
    body(doc, "En la parte inferior del sitio (fondo oscuro) encontrarás las 17 referencias "
              "bibliográficas que respaldan todo el contenido del sitio.")

    heading(doc, "7.1 Filtrar por Subtema", level=2)
    body(doc, "Cinco botones en la parte superior permiten filtrar las referencias:")
    styled_table(doc,
        ["Botón", "Muestra"],
        [
            ["Todas (17)",             "Todas las referencias bibliográficas"],
            ["Explícito vs Tácito",    "Solo fuentes del Subtema 1"],
            ["Innovación y Cultura",   "Solo fuentes del Subtema 2"],
            ["Trabajo Colaborativo",   "Solo fuentes del Subtema 3"],
            ["Aprendizaje",            "Solo fuentes del Subtema 4"],
        ],
        col_widths=[6, 10]
    )

    heading(doc, "7.2 Tarjetas de Referencias", level=2)
    body(doc, "Cada tarjeta muestra:")
    for item in [
        "#Número de referencia (esquina superior izquierda).",
        "Tipo de fuente: Libro (azul), Artículo (verde) o Informe (ámbar).",
        "Subtemas a los que aplica la fuente.",
        "Título completo de la obra, autores y año.",
        "Datos adicionales: editorial, ISBN, revista, volumen, páginas o DOI según corresponda.",
    ]:
        bullet(doc, item)

    heading(doc, "7.3 Acceder a las Fuentes", level=2)
    for item in [
        ('Botón verde "Ver artículo (DOI)": ',
         "Disponible para artículos científicos. Te dirige al artículo original vía DOI (identificador permanente)."),
        ('Botón azul "Ver fuente": ',
         "Disponible para libros e informes. Te dirige al catálogo Open Library, Harvard Business Review o McKinsey según la fuente."),
    ]:
        bullet(doc, item[1], bold_prefix=item[0])
    body(doc, "Todos los enlaces se abren en una nueva pestaña para no perder tu posición en el sitio.", color=GRIS_CLARO)
    divider(doc)

    # 8. Pie de página
    heading(doc, "8. Pie de Página (Footer)")
    styled_table(doc,
        ["Columna", "Contenido"],
        [
            ["Gestión del Conocimiento", "Descripción breve del proyecto"],
            ["Subtemas",                 "Links rápidos a las cuatro secciones"],
            ["Recursos",                 "Links a la bibliografía y tres fuentes clave externas"],
            ["Integrantes",              "Lista completa del equipo de trabajo"],
        ],
        col_widths=[6, 10]
    )
    divider(doc)

    # 9. Teclado
    heading(doc, "9. Navegación con el Teclado")
    styled_table(doc,
        ["Tecla", "Acción"],
        [
            ["Tab",             "Avanza al siguiente elemento interactivo"],
            ["Shift + Tab",     "Retrocede al elemento anterior"],
            ["Enter / Space",   "Activa el elemento seleccionado (links, botones)"],
            ["↑ / ↓",           "Desplazamiento vertical de la página"],
        ],
        col_widths=[5, 11]
    )
    divider(doc)

    # 10. Móvil
    heading(doc, "10. Uso en Dispositivos Móviles")
    body(doc, "El sitio está diseñado con enfoque mobile-first y se adapta a cualquier tamaño de pantalla:")
    for item in [
        "Menú hamburger para navegación en pantallas pequeñas.",
        "Estadísticas, bloques de contenido y videos pasan de múltiples columnas a 1 columna en móvil.",
        "La bibliografía pasa de 2 columnas a 1 columna en móvil.",
    ]:
        bullet(doc, item)
    body(doc, "Para una experiencia óptima de video en móvil, rota el dispositivo a orientación horizontal (modo paisaje).",
         color=GRIS_CLARO)
    divider(doc)

    # 11. FAQ
    heading(doc, "11. Preguntas Frecuentes")
    faqs = [
        ("¿El sitio requiere conexión a internet?",
         "Sí. Los videos de YouTube requieren conexión activa para cargar y reproducirse."),
        ("¿Por qué algunos videos tardan en cargar?",
         "Los videos usan carga diferida (lazy loading). El reproductor cargará al llegar a esa sección de la página."),
        ("¿Puedo descargar el contenido?",
         "El texto puede copiarse para uso educativo. Los videos deben reproducirse desde YouTube."),
        ("¿Las fuentes bibliográficas son verificables?",
         "Sí. Cada fuente tiene un botón que dirige al artículo original (DOI) o al registro del libro (Open Library). Los DOI son identificadores permanentes internacionales."),
        ("¿El sitio funciona sin JavaScript?",
         "No. Al ser una aplicación React (Next.js), requiere JavaScript habilitado en el navegador."),
        ("¿Cómo vuelvo al inicio?",
         "Haz clic en el logo 🧩 en la esquina superior izquierda, o selecciona \"Inicio\" en el menú."),
    ]
    for pregunta, respuesta in faqs:
        heading(doc, pregunta, level=3, space_before=8, color=AZUL_MEDIO)
        body(doc, respuesta, space_after=4)

    # Pie
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("Manual de usuario — versión 1.0 · mayo 2026")
    run.font.size  = Pt(9)
    run.font.color.rgb = GRIS_CLARO
    run.italic = True

    doc.save(r"d:\Archivos\Escritorio\GESTION DEL CONOCIMIENTO\Trabajo Final\docs\manual-de-usuario.docx")
    print("OK: manual-de-usuario.docx generado")


if __name__ == "__main__":
    crear_doc_tecnico()
    crear_manual_usuario()
    print("\nAmbos documentos generados en la carpeta docs/")

